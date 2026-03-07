import streamlit as st
import pandas as pd
from fpdf import FPDF
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

# --- Page Configuration ---
st.set_page_config(page_title="PF Interest Calculation", layout="wide")

# --- CSS for Tables & Layout ---
st.markdown("""
<style>
    .stDataFrame {font-size: 14px;}
    .summary-box { background-color: #f9f9f9; padding: 15px; border-radius: 10px; border: 1px solid #ddd; text-align: center; margin-bottom: 10px; }
    .total-box { background-color: #e6ffe6; padding: 15px; border-radius: 10px; border: 1px solid #b3ffb3; text-align: center; margin-bottom: 10px; }
    .stButton button { width: 100%; font-weight: bold; height: 50px; }
</style>
""", unsafe_allow_html=True)

st.title("📄 PF Interest Calculation Sheet By Pritam Konar")

# --- Inputs ---
col_h1, col_h2 = st.columns([2, 1])
with col_h1:
    school_name = st.text_input("School Name", placeholder="Enter School Name")
    employee_name = st.text_input("Employee Name", placeholder="Enter Employee Name")
with col_h2:
    start_year = st.number_input("Financial Year Start (e.g., 2024)", value=2024, step=1)
    
    # Store the default rate in session state to detect changes
    if 'prev_default_rate' not in st.session_state:
        st.session_state.prev_default_rate = 7.1
        
    rate_input = st.number_input("Default Rate of Interest (%)", min_value=0.0, value=7.1, step=0.1, format="%.2f")

# --- Sidebar ---
st.sidebar.header("Account Setup")
opening_balance_input = st.sidebar.number_input("Opening Balance (April 1st)", min_value=0.0, value=0.0, step=100.0, format="%.2f")

st.sidebar.header("Historical Adjustments")
st.sidebar.info("Use this to match old ledgers that contain manual typos or missing arrears.")
override_interest = st.sidebar.number_input("Override Final Interest (Enter 0 to use standard math)", value=0.0, format="%.2f")
round_1_dec = st.sidebar.checkbox("Round Total Interest to 1 Decimal", value=False)

# --- Helper: Months ---
def get_fy_months(start_year):
    m_names = ["APRIL", "MAY", "JUNE", "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER", "JANUARY", "FEBRUARY", "MARCH"]
    return [f"{m} {str(start_year if i < 9 else start_year + 1)[-2:]}" for i, m in enumerate(m_names)]

months_list = get_fy_months(start_year)

# --- Data Entry ---
# Initialize or update the Rate column if the global Default Rate changes
if 'input_data' not in st.session_state:
    st.session_state.input_data = pd.DataFrame({
        "Month": months_list, 
        "Dep_Before_15": [0.0]*12, "PFLR_Before_15": [0.0]*12, 
        "Dep_After_15": [0.0]*12, "PFLR_After_15": [0.0]*12, 
        "Withdrawal": [0.0]*12, 
        "Monthly_Rate": [rate_input]*12, # NEW COLUMN
        "Remarks": [""]*12
    })
else:
    st.session_state.input_data["Month"] = months_list
    # Auto-update all monthly rates if the user changes the Top Default Rate box
    if rate_input != st.session_state.prev_default_rate:
        st.session_state.input_data["Monthly_Rate"] = rate_input
        st.session_state.prev_default_rate = rate_input

st.subheader("Monthly Data Entry")
edited_df = st.data_editor(
    st.session_state.input_data,
    column_config={
        "Month": st.column_config.TextColumn("Month", disabled=True),
        "Dep_Before_15": st.column_config.NumberColumn("Deposit (< 15th)", format="₹ %.2f"),
        "PFLR_Before_15": st.column_config.NumberColumn("P.F.L.R (< 15th)", format="₹ %.2f"),
        "Dep_After_15": st.column_config.NumberColumn("Deposit (> 15th)", format="₹ %.2f"),
        "PFLR_After_15": st.column_config.NumberColumn("P.F.L.R (> 15th)", format="₹ %.2f"),
        "Withdrawal": st.column_config.NumberColumn("Withdrawal", format="₹ %.2f"),
        "Monthly_Rate": st.column_config.NumberColumn("Rate %", format="%.2f %%"), # EXPOSED IN UI
        "Remarks": st.column_config.TextColumn("Remarks")
    },
    hide_index=True, use_container_width=True, num_rows="fixed"
)

# --- Calculation Logic ---
def calculate_ledger(opening_bal, input_df, force_round, override_int):
    results = []
    current_bal = opening_bal
    total_interest = 0.0
    sum_dep_b = 0; sum_pflr_b = 0; sum_dep_a = 0; sum_pflr_a = 0; sum_with = 0

    for index, row in input_df.iterrows():
        month = row['Month']
        dep_before = row['Dep_Before_15']; pflr_before = row['PFLR_Before_15']
        dep_after = row['Dep_After_15']; pflr_after = row['PFLR_After_15']
        withdrawal = row['Withdrawal']; remarks = row['Remarks']
        
        # Pull the specific rate for this month
        month_rate = row['Monthly_Rate']

        sum_dep_b += dep_before; sum_pflr_b += pflr_before
        sum_dep_a += dep_after; sum_pflr_a += pflr_after; sum_with += withdrawal

        effective_deposit_for_interest = dep_before + pflr_before
        lowest_bal = max(0, current_bal + effective_deposit_for_interest - withdrawal)

        # Interest calculation using the CUSTOM MONTHLY RATE
        raw_interest = (lowest_bal * month_rate) / 1200
        interest = int(raw_interest * 100) / 100.0

        closing_bal = current_bal + dep_before + dep_after + pflr_before + pflr_after - withdrawal

        results.append({
            "Month": month, "Opening Balance": current_bal, "Dep (<15th)": dep_before, 
            "PFLR (<15th)": pflr_before, "Dep (>15th)": dep_after, "PFLR (>15th)": pflr_after, 
            "Withdrawal": withdrawal, "Lowest Balance": lowest_bal, 
            "Rate (%)": month_rate, # Store for display
            "Interest": interest, "Closing Balance": closing_bal, "Remarks": remarks
        })

        current_bal = closing_bal
        total_interest += interest
        
    # APPLY SPECIAL RULES
    if force_round:
        total_interest = round(total_interest, 1)
        
    # OVERRIDE RULE
    if override_int > 0:
        total_interest = override_int
    
    df_res = pd.DataFrame(results)
    totals = {
        "Dep (<15th)": sum_dep_b, "PFLR (<15th)": sum_pflr_b, "Dep (>15th)": sum_dep_a, 
        "PFLR (>15th)": sum_pflr_a, "Withdrawal": sum_with, "Interest": total_interest
    }
    return df_res, totals, current_bal

result_df, totals, final_principal = calculate_ledger(opening_balance_input, edited_df, round_1_dec, override_interest)

# --- Display Results ---
st.divider()
st.subheader("Calculation Result")

display_df = result_df.copy()
# Drop Rate column from the UI final calculation view to keep it clean like before
display_df_view = display_df.drop(columns=['Rate (%)'])

total_row = pd.DataFrame([{
    "Month": "TOTAL", "Opening Balance": None, "Dep (<15th)": totals["Dep (<15th)"], 
    "PFLR (<15th)": totals["PFLR (<15th)"], "Dep (>15th)": totals["Dep (>15th)"], 
    "PFLR (>15th)": totals["PFLR (>15th)"], "Withdrawal": totals["Withdrawal"], 
    "Lowest Balance": None, "Interest": totals["Interest"], "Closing Balance": None, "Remarks": ""
}])
display_df_view = pd.concat([display_df_view, total_row], ignore_index=True)
st.dataframe(display_df_view, use_container_width=True, hide_index=True)

# --- SUMMARY SECTION ---
final_total_balance = final_principal + totals['Interest']

st.markdown("### Final Summary")
col_s1, col_s2, col_s3 = st.columns(3)
with col_s1:
    st.markdown(f'<div class="summary-box"><div style="font-weight:bold; color:#555;">Principal</div><div style="font-size:22px;">₹ {final_principal:,.2f}</div></div>', unsafe_allow_html=True)
with col_s2:
    st.markdown(f'<div class="summary-box"><div style="font-weight:bold; color:#555;">Interest</div><div style="font-size:22px;">₹ {totals["Interest"]:,.2f}</div></div>', unsafe_allow_html=True)
with col_s3:
    st.markdown(f'<div class="total-box"><div style="font-weight:bold; color:#005500;">TOTAL</div><div style="font-size:22px; color:#005500;">₹ {final_total_balance:,.2f}</div></div>', unsafe_allow_html=True)

# --- PDF GENERATION ---
def create_pdf(df, school, name, year, rate, totals, final_bal):
    pdf = FPDF('L', 'mm', 'A4')
    pdf.set_margins(10, 10, 10)
    pdf.add_page()
    
    pdf.set_font('Arial', '', 14)
    pdf.multi_cell(0, 8, f"SCHOOL NAME :- {school}", 0, 'C')
    pdf.set_font('Arial', '', 9)
    pdf.cell(0, 8, f"INTEREST CALCULATION OF PROVIDENT FUND ACCOUNT FOR THE YEAR - {year}-{year+1}", 0, 1, 'C')
    pdf.ln(2) 
    pdf.set_font('Arial', '', 10)
    pdf.cell(140, 8, f"NAME :- {name}", 0, 0, 'L')
    pdf.cell(0, 8, f"RATE OF INTEREST:- {rate} %", 0, 1, 'R')
    pdf.ln(5) 

    w = {"mo": 26, "op": 26, "d1": 24, "p1": 24, "d2": 24, "p2": 24, "wi": 22, "lo": 26, "in": 20, "cl": 26, "re": 30}
    pdf.set_font('Arial', '', 8) 
    x, y = pdf.get_x(), pdf.get_y()
    
    pdf.rect(x, y, w['mo'], 12); pdf.set_xy(x, y+4); pdf.cell(w['mo'], 4, "Month", 0, 0, 'C')
    pdf.rect(x+w['mo'], y, w['op'], 12); pdf.set_xy(x+w['mo'], y+4); pdf.cell(w['op'], 4, "Opening Balance", 0, 0, 'C')
    
    grp1_x = x + w['mo'] + w['op']
    pdf.rect(grp1_x, y, w['d1']+w['p1'], 6); pdf.set_xy(grp1_x, y+1); pdf.cell(w['d1']+w['p1'], 4, "Deposit up to 15th day", 0, 0, 'C')
    pdf.rect(grp1_x, y+6, w['d1'], 6); pdf.set_xy(grp1_x, y+7); pdf.cell(w['d1'], 4, "Deposit", 0, 0, 'C')
    pdf.rect(grp1_x+w['d1'], y+6, w['p1'], 6); pdf.set_xy(grp1_x+w['d1'], y+7); pdf.cell(w['p1'], 4, "P.F.L.R", 0, 0, 'C')

    grp2_x = grp1_x + w['d1'] + w['p1']; grp2_w = w['d2'] + w['p2']
    pdf.rect(grp2_x, y, grp2_w, 6); pdf.set_xy(grp2_x, y); pdf.multi_cell(grp2_w, 3, "Deposit between\n16th & last day", 0, 'C')
    pdf.rect(grp2_x, y+6, w['d2'], 6); pdf.set_xy(grp2_x, y+7); pdf.cell(w['d2'], 4, "Deposit", 0, 0, 'C')
    pdf.rect(grp2_x+w['d2'], y+6, w['p2'], 6); pdf.set_xy(grp2_x+w['d2'], y+7); pdf.cell(w['p2'], 4, "P.F.L.R", 0, 0, 'C')

    curr_x = grp2_x + w['d2'] + w['p2']
    pdf.rect(curr_x, y, w['wi'], 12); pdf.set_xy(curr_x, y+4); pdf.cell(w['wi'], 4, "Withdrawal", 0, 0, 'C')
    curr_x += w['wi']
    pdf.rect(curr_x, y, w['lo'], 12); pdf.set_xy(curr_x, y+4); pdf.cell(w['lo'], 4, "Lowest Balance", 0, 0, 'C')
    curr_x += w['lo']
    pdf.rect(curr_x, y, w['in'], 12); pdf.set_xy(curr_x, y+2); pdf.multi_cell(w['in'], 4, "Interest\nfor month", 0, 'C')
    curr_x += w['in']
    pdf.rect(curr_x, y, w['cl'], 12); pdf.set_xy(curr_x, y+4); pdf.cell(w['cl'], 4, "Closing Balance", 0, 0, 'C')
    curr_x += w['cl']
    pdf.rect(curr_x, y, w['re'], 12); pdf.set_xy(curr_x, y+4); pdf.cell(w['re'], 4, "Remarks", 0, 0, 'C')

    pdf.set_xy(x, y + 12)
    pdf.set_font('Arial', '', 9) 
    row_h = 8
    def cell_c(w, txt, border=1): pdf.cell(w, row_h, str(txt), border, 0, 'C')

    for _, row in df.iterrows():
        cell_c(w['mo'], row['Month']); cell_c(w['op'], f"{row['Opening Balance']:.2f}"); cell_c(w['d1'], f"{row['Dep (<15th)']:.2f}"); cell_c(w['p1'], f"{row['PFLR (<15th)']:.2f}"); cell_c(w['d2'], f"{row['Dep (>15th)']:.2f}"); cell_c(w['p2'], f"{row['PFLR (>15th)']:.2f}"); cell_c(w['wi'], f"{row['Withdrawal']:.2f}"); cell_c(w['lo'], f"{row['Lowest Balance']:.2f}"); cell_c(w['in'], f"{row['Interest']:.2f}"); cell_c(w['cl'], f"{row['Closing Balance']:.2f}"); cell_c(w['re'], row['Remarks'])
        pdf.ln()

    pdf.set_font('Arial', '', 9)
    cell_c(w['mo'], "Total"); cell_c(w['op'], ""); cell_c(w['d1'], f"{totals['Dep (<15th)']:.2f}"); cell_c(w['p1'], f"{totals['PFLR (<15th)']:.2f}"); cell_c(w['d2'], f"{totals['Dep (>15th)']:.2f}"); cell_c(w['p2'], f"{totals['PFLR (>15th)']:.2f}"); cell_c(w['wi'], f"{totals['Withdrawal']:.2f}"); cell_c(w['lo'], ""); cell_c(w['in'], f"{totals['Interest']:.2f}"); cell_c(w['cl'], ""); cell_c(w['re'], "")
    pdf.ln(12)

    pdf.ln(5)
    pdf.set_font('Arial', '', 10)
    pdf.cell(30, 6, "Principal", 0, 0); pdf.cell(30, 6, f": {final_bal:.2f}", 0, 1)
    pdf.cell(30, 6, "Interest", 0, 0); pdf.cell(30, 6, f": {totals['Interest']:.2f}", 0, 1)
    pdf.cell(30, 6, "TOTAL", 0, 0); pdf.cell(30, 6, f": {(final_bal + totals['Interest']):.2f}", 0, 0)
    pdf.set_x(230); pdf.cell(40, 6, "Signature of HM", 0, 1, 'C')

    return pdf.output(dest='S').encode('latin-1')

# --- EXCEL GENERATION (Using openpyxl to match exact template) ---
def create_excel(df, school, name, year, rate, totals, final_bal):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # ── Page setup ──
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left   = 0.25
    ws.page_margins.right  = 0.25
    ws.page_margins.top    = 0.25
    ws.page_margins.bottom = 0.25
    ws.page_margins.header = 0.0
    ws.page_margins.footer = 0.0
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered   = True

    # ── Column widths ──
    col_widths = {
        "A": 15.42578125, "B": 14.0, "C": 11.42578125, "D": 10.5703125,
        "E": 11.42578125, "F": 9.7109375, "G": 12.5703125, "H": 12.7109375,
        "I": 11.85546875, "J": 12.5703125, "K": 10.7109375, "L": 9.140625,
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # ── Row heights ──
    row_heights = {
        1: 21.0, 2: 21.0, 3: 21.0, 4: 21.0, 5: 21.0, 6: 21.0,
        7: 21.95, 8: 21.95, 9: 21.95, 10: 21.95, 11: 21.95, 12: 21.95,
        13: 21.95, 14: 21.95, 15: 21.95, 16: 21.95, 17: 21.95, 18: 21.95,
        19: 21.95, 20: 21.95, 21: 21.95, 22: 21.95, 23: 110.1,
    }
    for row, height in row_heights.items():
        ws.row_dimensions[row].height = height

    # ── Merged cells ──
    merges = [
        "A1:K1", "A2:K2", "A3:G3", "H3:K3",
        "A4:A6", "B4:B6", "C4:D5", "E4:F5",
        "G4:G6", "H4:H6", "I4:I6", "J4:J6", "K4:K6",
        "C22:D22", "A23:K23",
    ]
    for m in merges:
        ws.merge_cells(m)

    # ── Style helpers ──
    def thin(): return Side(border_style="thin")
    def all_thin(): return Border(left=thin(), right=thin(), top=thin(), bottom=thin())
    def font_cambria9(): return Font(name="Cambria", size=9)
    def font_calibri10(): return Font(name="Calibri", size=10)
    def align_cc(wrap=False): return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

    # ── Row 1 ──
    ws["A1"].value = school if school else "SHYAMBAZAR D.N. HIGH SCHOOL"
    ws["A1"].font = Font(name="Calibri", size=11)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A1"].border = Border(left=thin(), right=thin(), top=thin(), bottom=thin())
    for col in "BCDEFGHIJ": ws[f"{col}1"].border = Border(top=thin(), bottom=thin())
    ws["K1"].border = Border(right=thin(), top=thin(), bottom=thin())

    # ── Row 2 ──
    ws["A2"].value = f"                                   INTEREST CALCULATION OF PROVIDENT FUND ACCOUNT FOR THE YEAR  {year}-{year+1}"
    ws["A2"].font = font_cambria9()
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    for col in "ABCDEFGHIJK":
        ws[f"{col}2"].border = Border(left=thin() if col == "A" else None, right=thin() if col == "K" else None, top=thin(), bottom=thin())

    # ── Row 3 ──
    n = name if name else ""
    ws["A3"].value = f"NAME :- {n}   (School Code-92)                     "
    ws["A3"].font = font_cambria9()
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")
    ws["H3"].value = f"RATE OF INTEREST:- {rate} %"
    ws["H3"].font = font_cambria9()
    ws["H3"].alignment = Alignment(horizontal="center", vertical="center")
    for col in "ABCDEFG": ws[f"{col}3"].border = Border(left=thin() if col == "A" else None, right=thin() if col == "G" else None, top=thin(), bottom=thin())
    for col in "HIJK": ws[f"{col}3"].border = Border(left=thin() if col == "H" else None, right=thin() if col == "K" else None, top=thin(), bottom=thin())

    # ── Rows 4-6 Headers ──
    headers_r4 = {
        "A4": "Month", "B4": "Opening balance", "C4": "Deposit up to 15th day Deposit ",
        "E4": "Depisit between16th & last day", "G4": "Withdrawal", "H4": "Lowest Balance",
        "I4": "Interest for the month", "J4": "Closing Balance", "K4": "Remarks",
    }
    for cell_addr, text in headers_r4.items():
        ws[cell_addr].value = text
        ws[cell_addr].font = font_cambria9()
        ws[cell_addr].alignment = align_cc(wrap=True)

    ws["C6"].value = "Deposit"
    ws["D6"].value = "P.F.L.R"
    ws["E6"].value = "Deposit"
    ws["F6"].value = "P.F.L.R"
    for col in "CDEF":
        ws[f"{col}6"].font = font_cambria9()
        ws[f"{col}6"].alignment = align_cc(wrap=True)

    border_map_4 = {"A4": all_thin(), "B4": all_thin(), "C4": all_thin(), "D4": Border(right=thin(), top=thin()), "E4": all_thin(), "F4": Border(right=thin(), top=thin()), "G4": all_thin(), "H4": all_thin(), "I4": all_thin(), "J4": all_thin(), "K4": all_thin()}
    for addr, b in border_map_4.items(): ws[addr].border = b
    border_map_5 = {"A5": Border(left=thin(), right=thin()), "B5": Border(left=thin(), right=thin()), "C5": Border(left=thin(), bottom=thin()), "D5": Border(right=thin(), bottom=thin()), "E5": Border(left=thin(), bottom=thin()), "F5": Border(right=thin(), bottom=thin()), "G5": Border(left=thin(), right=thin()), "H5": Border(left=thin(), right=thin()), "I5": Border(left=thin(), right=thin()), "J5": Border(left=thin(), right=thin()), "K5": Border(left=thin(), right=thin())}
    for addr, b in border_map_5.items(): ws[addr].border = b
    for col in "ABCDEFGHIJK": ws[f"{col}6"].border = Border(left=thin() if col in "ABCDEFGHIJK" else None, right=thin() if col in "ABCDEFGHIJK" else None, top=thin() if col in "CDEF" else None, bottom=thin())

    # ── Rows 7-18 Data ──
    for idx, row_data in df.iterrows():
        r = 7 + idx
        
        # Clean Month Name
        month_str = str(row_data['Month']).split(" ")[0].upper()
        if month_str == "MAY": month_str = "May"
        elif month_str in ["JUNE", "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER", "JANUARY"]:
            month_str += " "

        ws[f"A{r}"].value = month_str
        ws[f"B{r}"].value = row_data['Opening Balance'] if row_data['Opening Balance'] != 0 else ""
        ws[f"C{r}"].value = row_data['Dep (<15th)'] if row_data['Dep (<15th)'] != 0 else ""
        ws[f"D{r}"].value = row_data['PFLR (<15th)'] if row_data['PFLR (<15th)'] != 0 else ""
        ws[f"E{r}"].value = row_data['Dep (>15th)'] if row_data['Dep (>15th)'] != 0 else ""
        ws[f"F{r}"].value = row_data['PFLR (>15th)'] if row_data['PFLR (>15th)'] != 0 else ""
        ws[f"G{r}"].value = row_data['Withdrawal'] if row_data['Withdrawal'] != 0 else ""
        ws[f"H{r}"].value = row_data['Lowest Balance'] if row_data['Lowest Balance'] != 0 else ""
        ws[f"I{r}"].value = row_data['Interest'] if row_data['Interest'] != 0 else ""
        ws[f"J{r}"].value = row_data['Closing Balance'] if row_data['Closing Balance'] != 0 else ""
        
        remarks = str(row_data['Remarks']).strip()
        if remarks:
            ws[f"K{r}"].value = remarks
        else:
            # Uses the custom rate specific to that month for the excel file printing!
            ws[f"K{r}"].value = row_data['Rate (%)'] / 100.0 

    # Format Rows 7-22
    for r in range(7, 23):
        for col in "ABCDEFGHIJK":
            cell = ws[f"{col}{r}"]
            if col == "A":
                cell.font = font_calibri10()
                cell.alignment = align_cc()
            else:
                cell.font = font_calibri10()
                cell.alignment = align_cc()
                
            if col == "K" and r <= 18:
                cell.font = font_cambria9()
                cell.alignment = align_cc(wrap=True)
                if isinstance(cell.value, float):
                    cell.number_format = "0.00%"

            if r == 22 and col == "D":
                cell.border = Border(right=thin(), top=thin(), bottom=thin())
            else:
                cell.border = all_thin()

    # Total Row (19)
    ws["A19"].value = "Total"
    ws["C19"].value = totals['Dep (<15th)'] if totals['Dep (<15th)'] != 0 else ""
    ws["D19"].value = totals['PFLR (<15th)'] if totals['PFLR (<15th)'] != 0 else ""
    ws["E19"].value = totals['Dep (>15th)'] if totals['Dep (>15th)'] != 0 else ""
    ws["F19"].value = totals['PFLR (>15th)'] if totals['PFLR (>15th)'] != 0 else ""
    ws["G19"].value = totals['Withdrawal'] if totals['Withdrawal'] != 0 else ""
    ws["I19"].value = totals['Interest'] if totals['Interest'] != 0 else ""

    # Principal Row (20)
    ws["A20"].value = "Principal"
    ws["B20"].value = final_bal

    # Interest Row (21)
    ws["A21"].value = "Interest"
    ws["B21"].value = totals['Interest']

    # Final Total Row (22)
    ws["A22"].value = "TOTAL"
    ws["B22"].value = final_bal + totals['Interest']
    
    # ── Row 23 ──
    ws["A23"].value = "Sisnature of HM"
    ws["A23"].font = font_calibri10()
    ws["A23"].alignment = Alignment(horizontal="right", vertical="bottom")
    ws["A23"].border = all_thin()
    for col in "BCDEFGHIJ": ws[f"{col}23"].border = Border(top=thin(), bottom=thin())
    ws["K23"].border = Border(right=thin(), top=thin(), bottom=thin())

    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

# --- DOWNLOAD BUTTONS ---
st.write("") 
col_d1, col_d2 = st.columns(2)

pdf_bytes = create_pdf(result_df, school_name, employee_name, start_year, rate_input, totals, final_principal)
with col_d1:
    st.download_button("📄 Download PDF", pdf_bytes, f"PF_{start_year}.pdf", 'application/pdf', use_container_width=True)

excel_bytes = create_excel(result_df, school_name, employee_name, start_year, rate_input, totals, final_principal)
with col_d2:
    st.download_button("📊 Download Excel", excel_bytes, f"PF_{start_year}.xlsx", 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', use_container_width=True)
